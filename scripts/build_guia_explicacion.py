"""Guía de explicación para quien gestiona el proyecto y no programa.

No altera el motor ni los resultados. Genera PDF y, si está instalado, DOCX.
"""

from __future__ import annotations

from pathlib import Path

from fpdf import FPDF

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
FONT = Path(r"C:\Windows\Fonts\arial.ttf")
FONTB = Path(r"C:\Windows\Fonts\arialbd.ttf")
FONTI = Path(r"C:\Windows\Fonts\ariali.ttf")


class Guide(FPDF):
    def __init__(self) -> None:
        super().__init__(format="A4", unit="mm")
        self.set_auto_page_break(auto=True, margin=16)
        self.add_font("Arial", "", str(FONT))
        self.add_font("Arial", "B", str(FONTB))
        self.add_font("Arial", "I", str(FONTI))
        self.blocks: list[tuple[str, str]] = []

    def footer(self) -> None:
        if self.page_no() == 1:
            return
        self.set_y(-12)
        self.set_font("Arial", "", 8)
        self.set_text_color(80, 80, 80)
        self.cell(0, 6, f"Guía de explicación del proyecto  ·  {self.page_no()}", align="C")
        self.set_text_color(0, 0, 0)

    def h1(self, text: str) -> None:
        self.set_x(self.l_margin)
        self.ln(1.5)
        self.set_font("Arial", "B", 13)
        self.set_text_color(20, 40, 70)
        self.multi_cell(0, 6.5, text)
        self.set_text_color(0, 0, 0)
        self.ln(0.6)
        self.blocks.append(("h1", text))

    def p(self, text: str) -> None:
        self.set_x(self.l_margin)
        self.set_font("Arial", "", 11)
        self.multi_cell(0, 5.3, text)
        self.ln(1.1)
        self.blocks.append(("p", text))

    def bullet(self, text: str) -> None:
        self.set_x(self.l_margin)
        self.set_font("Arial", "", 11)
        self.multi_cell(0, 5.3, "•  " + text)
        self.ln(0.3)
        self.blocks.append(("b", text))

    def term(self, name: str, text: str) -> None:
        self.set_x(self.l_margin)
        self.set_font("Arial", "B", 11)
        self.multi_cell(0, 5.2, name)
        self.set_x(self.l_margin)
        self.set_font("Arial", "", 10)
        self.multi_cell(0, 4.8, text)
        self.ln(0.8)
        self.blocks.append(("term", name + "\n" + text))

    def qa(self, q: str, a: str) -> None:
        self.set_x(self.l_margin)
        self.set_font("Arial", "B", 11)
        self.multi_cell(0, 5.2, q)
        self.set_x(self.l_margin)
        self.set_font("Arial", "", 11)
        self.multi_cell(0, 5.2, a)
        self.ln(1.0)
        self.blocks.append(("qa", q + "\n" + a))


def build() -> Guide:
    pdf = Guide()
    pdf.add_page()
    pdf.ln(22)
    pdf.set_font("Arial", "B", 18)
    pdf.multi_cell(0, 8, "Cómo explicar este proyecto", align="C")
    pdf.ln(4)
    pdf.set_font("Arial", "", 12)
    pdf.multi_cell(
        0,
        6,
        "Guía de estudio para quien ha seguido el trabajo\ny no necesita programar para contarlo.",
        align="C",
    )
    pdf.ln(8)
    pdf.set_font("Arial", "I", 11)
    pdf.multi_cell(
        0,
        5.5,
        "Simulación de la adsorción de azul de metileno\nsobre óxido de grafeno y carbón activado.",
        align="C",
    )
    pdf.ln(10)
    pdf.set_font("Arial", "", 10)
    pdf.multi_cell(
        0,
        5,
        "Complementa la memoria técnica. No añade resultados\ny no sustituye un experimento de laboratorio.",
        align="C",
    )

    pdf.h1("1. Si me preguntan qué es este proyecto, ¿qué digo?")
    pdf.p(
        "Texto para aprender casi de memoria: Hemos simulado, en un plano, cómo el azul de metileno puede quedar retenido por dos materiales: óxido de grafeno, GO, y carbón activado, AC. No es una simulación de moléculas. Son objetos que se mueven al azar, se tocan según la distancia y, si cabe su masa, el colorante queda adsorbido. Queríamos ver, con las mismas reglas, cuánto retiene cada material y cuánto varía eso al repetir la simulación. Con 40 repeticiones por material, el GO retuvo de media unos 187 objetos de 200, cerca del 94 % y unos 375 mg por gramo de sólido. El AC retuvo unos 133, cerca del 67 % y unos 266 mg/g. Eso describe este modelo. No demuestra que el GO sea mejor en cualquier experimento real."
    )

    pdf.h1("2. La misma idea, contada como una historia")
    pdf.p(
        "Imagina agua con un colorante azul, el azul de metileno. En el modelo hay 4 miligramos de colorante en 40 mililitros, a 25 grados y pH 6. Ese cociente es C0 = 4 mg / 0,04 L = 100 mg/L: es una magnitud derivada, no un dato bibliográfico aparte. Introducimos 10 miligramos de GO o, en otra simulación, 10 miligramos de AC. El programa no dibuja el agua molécula a molécula. Representa el colorante como 200 objetos que dan saltos al azar dentro de un cuadrado. Si el sólido es GO, ese sólido también se mueve. Si es AC, se queda quieto. Cuando un objeto de colorante se acerca lo suficiente a una unidad de sólido, hay contacto. Entonces se mira si en esa unidad aún cabe la masa de un objeto entero, 0,02 miligramos. Si cabe, el colorante queda adsorbido y deja de moverse. Como el movimiento es aleatorio, una sola vez no basta: repetimos 40 veces cada material, con semillas distintas, y miramos la media y la dispersión."
    )

    pdf.h1("3. Qué hace realmente la simulación")
    steps = [
        "Se crea el sistema: cuadrado, masas, tamaños y la regla de adsorción.",
        "Se colocan 200 objetos de colorante y 100 unidades de sólido, al azar y sin solaparse.",
        "Los objetos de colorante que siguen libres dan un salto aleatorio.",
        "El GO también se mueve. El AC permanece fijo.",
        "Se mide la distancia. Si es menor o igual que la suma de radios, hay contacto.",
        "Si en esa unidad cabe un objeto entero, se adsorbe. En la campaña final esa regla es siempre sí, cuando hay sitio.",
        "Se cuentan los objetos adsorbidos y los contactos. No son el mismo número.",
        "Se calcula qt: masa adsorbida dividida por la masa del sólido, en mg/g.",
        "Se repite con otra semilla. La misma semilla vuelve a dar el mismo recorrido.",
        "Con las 40 repeticiones se calculan media, dispersión y el resto de resúmenes.",
    ]
    for i, text in enumerate(steps, 1):
        pdf.bullet(f"{i}. {text}")
    pdf.ln(1)
    pdf.p(
        "Ese orden es el del programa: primero mover, después detectar el contacto, después adsorber y anotar. La ventana de presentación usa el mismo recorrido. No hay un segundo cálculo escondido."
    )

    pdf.h1("4. Por qué GO y AC, y en qué se diferencian")
    pdf.p(
        "El proyecto compara dos adsorbentes frente al mismo colorante y con el mismo protocolo de ejecución. La diferencia que el programa sí representa es de dibujo y de movimiento, no una fotografía del material al microscopio."
    )
    pdf.term(
        "GO, óxido de grafeno",
        "Se mueve. En el dibujo final cada unidad es un cuadrado de 7×7 píxeles. Su capacidad de referencia, qmax, es 764,7 mg/g. En la campaña contacta mucho y retiene de media unos 187 objetos.",
    )
    pdf.term(
        "AC, carbón activado",
        "Se queda fijo. Cada unidad es un cuadrado de 2×2 píxeles. Su qmax de referencia es 412,2 mg/g. Contacta menos y retiene de media unos 133 objetos. En cambio, convierte en adsorción una fracción mayor de los contactos que tiene.",
    )
    pdf.p(
        "Esos tamaños son una decisión de representación. No hay que decir que el grafeno mide 7 píxeles en la realidad ni que el carbón mide 2."
    )

    pdf.h1("5. Qué significan los resultados")
    pdf.p(
        "De 200 objetos de colorante, el GO adsorbe de media 187,45. Eso es el 93,73 %. Como cada objeto pesa 0,02 mg y el sólido pesa 0,01 g, qt sale 374,9 mg/g. El AC adsorbe de media 133,175 objetos, el 66,59 %, y qt sale 266,35 mg/g. La diferencia es de unos 54,3 objetos y unos 108,5 mg/g."
    )
    pdf.p(
        "En este modelo, y solo en estas condiciones de cálculo, el GO produjo más adsorción que el AC. También tuvo más contactos, unos 606 frente a unos 321. Eso no autoriza a decir que el GO es experimentalmente mejor en cualquier situación, ni que el laboratorio vaya a medir esos mismos números."
    )
    pdf.p(
        "Una semilla de GO llega a 392 mg/g, un poco por encima del equilibrio de referencia del GO, 380,7 mg/g. No es un fallo que haya que corregir: el programa no está programado para frenar en ese valor. Adsorbe objetos enteros, y el azar de cada semilla mueve un poco el resultado. La dispersión es pequeña: el coeficiente de variación de los objetos adsorbidos es 2,5 % en GO y 4,2 % en AC."
    )

    pdf.h1("6. Langmuir y la curva cinética, sin fórmulas largas")
    pdf.term(
        "Langmuir",
        "Es una forma conocida de describir cuánto colorante cabría en equilibrio. qmax es el techo de esa descripción. KL dice con qué fuerza se asocia el colorante a esa curva. qe es la cantidad en equilibrio para nuestras condiciones: 380,7 mg/g en GO y 290,9 mg/g en AC. Ce es la concentración que quedaría en el líquido en ese equilibrio. El programa los tiene al lado, como referencia. No los usa para decidir cada contacto.",
    )
    pdf.term(
        "PSO, pseudo-segundo orden",
        "Es una curva que describe cómo iría creciendo la cantidad adsorbida a lo largo de minutos de laboratorio. El GO y el AC tienen curvas distintas. La curva de referencia del GO usa k2 = 0,0002. En otra tabla del repositorio el k2 del GO aparece como 0,001. Esa diferencia está en la información guardada y no se ha resuelto con una nueva consulta bibliográfica. La simulación avanza por pasos de cálculo. Esos pasos no se han convertido en minutos. Un intento de relacionarlos, llamado alfa, se exploró y no se adoptó.",
    )

    pdf.h1("7. Por qué se repite la simulación")
    pdf.p(
        "Si se tira un dado una sola vez, no se conoce su comportamiento. Si se tira muchas veces, aparece la variabilidad. Aquí pasa lo mismo: el salto de cada objeto es aleatorio. Una semilla es el número que arranca esa lotería. La misma semilla repite el mismo recorrido. Otra semilla es otra realización. La media resume las 40. La desviación dice cuánto se separan de esa media. El CV es esa desviación expresada como porcentaje de la media: sirve para comparar dispersiones. Reproducibilidad significa que, con la misma semilla y la misma configuración, el resultado se puede volver a obtener."
    )

    pdf.h1("8. Cómo se tomaron las decisiones")
    pdf.term("Definir el problema", "Qué queríamos: representar el encuentro entre colorante y sólido, y anotar cuánta masa queda retenida. Decidimos un modelo de objetos en un plano, no de moléculas.")
    pdf.term("Construir el motor", "Qué hicimos: separar movimiento, contacto y adsorción. Qué vimos: que se puede contar masa, contactos y objetos adsorbidos sin mezclarlos.")
    pdf.term("Sensibilidad", "Qué queríamos saber: si alargar los pasos o cambiar el tamaño del salto altera el resultado. Qué vimos: sí cambia, y hacia muchos pasos el crecimiento se frena, sin una meseta perfecta. No convertimos el salto en un coeficiente de difusión medido.")
    pdf.term("Geometría", "Primero el colorante era más grande en el dibujo y los dos sólidos se parecían. Después se fijó el dibujo final: colorante 1×1, GO 7×7 móvil, AC 2×2 fijo. La serie antigua se archivó. No se mezclan.")
    pdf.term("Tiempo y P_ads", "Se comparó la simulación con las curvas en minutos, sin fijar una equivalencia. El barrido de P_ads cambia lo adsorbido y también las trayectorias. En el GO hay configuraciones en las que la media no baja de forma monótona al bajar P_ads. No se eligió un valor experimental. La campaña final usa la regla 1 como decisión de cálculo: si hay contacto y cabe, se adsorbe.")
    pdf.term("Campaña final", "40 semillas de GO y 40 de AC, salto 1,5, 2000 pasos. No se reajustó nada a mitad. Esos son los números que hay que citar.")

    pdf.h1("9. Qué no debo afirmar")
    for line in [
        "No es dinámica molecular y no representa moléculas de agua.",
        "No demuestra con un experimento que el GO sea mejor en general.",
        "P_ads = 1 no es una probabilidad medida en el laboratorio. Es la regla de cálculo de la campaña.",
        "S* = 0,55 es un dato publicado para otro carbón y este colorante. No es la regla que usó la campaña.",
        "Sigma, el tamaño del salto, no debe presentarse como una propiedad física ya medida.",
        "Los pasos de la simulación no equivalen automáticamente a minutos.",
        "Langmuir no obliga al programa a terminar en qe.",
        "Los cuadrados de 1, 7 y 2 píxeles son dibujo, no tamaños reales.",
    ]:
        pdf.bullet(line)

    pdf.h1("10. Glosario corto")
    glossary = [
        ("MB", "Azul de metileno, el colorante. En el modelo son 200 objetos de 0,02 mg."),
        ("GO", "Óxido de grafeno. Adsorbente móvil, dibujado a 7×7."),
        ("AC", "Carbón activado. Adsorbente fijo, dibujado a 2×2."),
        ("Adsorción", "El colorante queda asociado al sólido y deja de estar libre."),
        ("Adsorbente", "El sólido que retiene: GO o AC. Hay 100 unidades y 10 mg en total."),
        ("Adsorbato", "Lo que se retiene: el azul de metileno."),
        ("qt", "Cantidad adsorbida por gramo de sólido, en mg/g. Es el número que se compara."),
        ("qe", "La cantidad en el equilibrio de referencia. No es el final obligatorio de la simulación."),
        ("qmax", "Techo de la descripción de Langmuir: 764,7 mg/g en GO y 412,2 mg/g en AC."),
        ("KL", "Parámetro de esa descripción. 0,206 para GO y 0,088 para AC. No es una probabilidad."),
        ("Ce", "Concentración que quedaría en el líquido en el equilibrio de referencia."),
        ("Langmuir", "Modelo de equilibrio usado como referencia, no como freno del programa."),
        ("PSO", "Curva de cómo crecería qt con los minutos de laboratorio. Cada material tiene la suya."),
        ("Movimiento browniano", "Aquí: saltos al azar. No es una trayectoria molecular calibrada."),
        ("Sigma", "Tamaño típico de ese salto, en píxeles por paso. En la campaña, 1,5."),
        ("Semilla", "Número que arranca el azar. La misma semilla repite el mismo resultado."),
        ("Simulación estocástica", "El resultado depende del azar. Por eso se repite y se mira la media."),
        ("Contacto", "Dos objetos están lo bastante cerca. No implica que haya adsorción."),
        ("P_ads", "Regla de cálculo después del contacto. En la campaña vale 1: si cabe, se adsorbe."),
        ("S*", "0,55, publicado para un carbón relacionado. No sustituye a P_ads."),
        ("CSV", "Tabla de datos. Se abre con Excel o un programa parecido."),
        ("JSON", "Ficha estructurada: configuración, serie y resultado de una corrida."),
        ("PNG", "Imagen. En este proyecto, un gráfico hecho a partir de los resultados."),
        ("Test", "Comprobación automática de que una parte del programa sigue bien. No es un experimento."),
        ("Script", "Programa auxiliar. La mayoría documenta, abre la ventana o comprueba el entorno. El empaquetado del ejecutable está en packaging/ y en scripts/build_windows.ps1."),
        ("Repositorio", "La carpeta del proyecto, con su historial de cambios."),
        ("Commit", "Una foto guardada de ese historial: qué cambió y cuándo. No es un resultado científico."),
        ("Reproducibilidad", "Misma configuración y misma semilla, mismo recorrido. Los archivos permiten volver a la procedencia de un gráfico."),
    ]
    for name, text in glossary:
        pdf.term(name, text)

    pdf.h1("11. Dónde está cada cosa")
    pdf.p("Detalle ampliado en docs/mapa_del_repositorio.md. Para explicarlo basta esto:")
    folders = [
        ("src/go_mb/", "El programa.", "Calcula la simulación.", "El motor vive aquí."),
        ("tests/", "135 funciones de prueba identificadas.", "Ver que nada se ha roto.", "No son experimentos y esta guía no las ha vuelto a ejecutar."),
        ("scripts/", "Herramientas.", "Tareas concretas, aparte del motor.", "No son la ciencia."),
        ("data/final_campaign/", "80 simulaciones y gráficos.", "Las medias que se citan.", "De aquí salen los números finales."),
        ("data/sensitivity/", "Exploraciones anteriores.", "Probar salto, pasos, dibujo y P_ads.", "No se mezcla con la campaña final."),
        ("data/analysis/", "Lecturas de lo ya guardado.", "Tiempo de referencia y protocolo.", "No es una campaña nueva."),
        ("docs/", "Textos.", "Estudiar y explicar.", "Memoria, guía breve y esta guía."),
        ("packaging/", "Empaquetado.", "Crear el programa de Windows.", "Para presentarlo sin el entorno de desarrollo."),
    ]
    for name, what, why, how in folders:
        pdf.term(name, f"{what} {why} Frase útil: {how}")
    pdf.p(
        "En la raíz están Iniciar_Simulacion.vbs y .bat, que abren la ventana; README.md, la puerta de entrada; requirements.txt, la lista de bibliotecas; y memoria.md, notas de trabajo. data/final_campaign_reported/ es una copia de consulta de la campaña final."
    )

    pdf.h1("12. Qué es tests/")
    pdf.p(
        "Un test es una comprobación automática: el programa se hace una pregunta concreta y verifica la respuesta. Por ejemplo, que la masa no desaparezca, que la misma semilla se repita, o que el contacto sea una distancia y no un sorteo. Existen para que un cambio posterior no rompa en silencio lo que ya funcionaba. El repositorio contiene 135 funciones de prueba identificadas en tests/. Esta guía no las ha vuelto a ejecutar. No son 135 experimentos científicos."
    )
    pdf.p(
        "Comprueban, entre otras cosas, el movimiento, el contacto, la adsorción, la geometría final, la conservación de masa, las campañas reducidas y que la ventana use el mismo cálculo que el modo sin ventana. Pausar no avanza el estado. Reiniciar con la misma semilla vuelve a empezar igual."
    )

    pdf.h1("13. Qué es scripts/")
    pdf.p(
        "Un script es un archivo que hace una tarea auxiliar y se deja fuera del motor para no mezclar la ciencia con la herramienta. No todo scripts/ sirve para empaquetar. packaging/ y scripts/build_windows.ps1 construyen el ejecutable de Windows, y ese archivo es solo de desarrollo. El resto son herramientas: build_documentation_pdf.py genera la memoria técnica; build_guia_explicacion.py genera esta guía; launch_view.ps1 abre la visualización; check_view_env.ps1 comprueba si esa ventana puede abrirse; y run_ac_sensitivity_batch.py se usó en caracterizaciones del carbón. Ese último no es la campaña final y no hace falta volver a ejecutarlo para explicar los resultados."
    )

    pdf.h1("14. CSV, JSON e imágenes")
    pdf.term("CSV", "Una tabla. En la campaña, per_run_records.csv tiene una fila por simulación y comparison_go_ac.csv tiene las medias. Se pueden abrir en una hoja de cálculo.")
    pdf.term("JSON", "Una ficha. Guarda la configuración, la semilla y la serie de una corrida, o el resumen estadístico. Sirve para saber de dónde salió un número.")
    pdf.term("PNG", "Un gráfico. Las figuras de la campaña final están en data/final_campaign/figures/. Muestran semillas, distribuciones, contactos y la comparación con las referencias. El eje de la simulación son pasos, no minutos.")

    pdf.h1("15. Por qué hay tantos archivos")
    pdf.p(
        "Una simulación que se quiere poder explicar guarda la configuración, cada semilla, el resultado individual, el agregado, las estadísticas, los gráficos y una nota de procedencia: qué protocolo se usó y qué no pretende ser. La idea es la trazabilidad. Si alguien pregunta de dónde salió un gráfico, se puede volver a la tabla y a la configuración que lo generaron. Por eso la campaña final no se mezcló con las carpetas de las pruebas anteriores."
    )

    pdf.h1("16. No hace falta programar para entender el código")
    pdf.p("Cada archivo importante se puede explicar con tres ideas: qué hace, por qué existe y qué frase basta.")
    files = [
        ("engine.py", "Ordena el paso: mover, tocar, adsorber, anotar.", "Es el corazón del cálculo.", "El ciclo siempre es ese."),
        ("motion.py", "Da el salto aleatorio y rebota en el borde.", "Separa el movimiento del resto.", "Sigma es el tamaño del salto, no un dato de laboratorio."),
        ("contact.py", "Decide si dos objetos están lo bastante cerca.", "El contacto no es un sorteo.", "Distancia menor o igual que la suma de radios."),
        ("adsorption.py", "Pasa el objeto al sólido si cabe su masa.", "La regla va después del contacto.", "En la campaña, si cabe, se adsorbe."),
        ("config.py", "Guarda masas, tamaños, qmax y condiciones.", "Para no esconder los números.", "25 °C, pH 6, 4 mg, 10 mg, 40 mL, 100 mg/L."),
        ("models.py", "Calcula Langmuir y la curva en minutos.", "Para tener la referencia al lado.", "No frena la simulación."),
        ("metrics.py", "Calcula qt y los recuentos.", "Para hablar de cantidades y no solo de dibujos.", "qt, contactos y objetos adsorbidos van por separado."),
        ("stats.py", "Media y dispersión de las semillas.", "Para no citar una sola tirada.", "Resume 40 corridas."),
        ("io.py", "Escribe CSV y JSON.", "Para guardar y poder volver atrás.", "La tabla y la ficha."),
        ("viz.py e interactive.py", "Dibujan. La ventana usa el mismo paso.", "Para verlo sin cambiar la física.", "Misma semilla, mismos números."),
    ]
    for name, does, why, phrase in files:
        if why.lower().startswith("para "):
            purpose = f"Existe para {why[5].lower()}{why[6:]}"
        else:
            purpose = why
        pdf.term(name, f"{does} {purpose} Frase: {phrase}")

    pdf.h1("17. Preguntas que pueden hacerme")
    pairs = [
        ("¿Por qué Python?", "Porque el programa está escrito en ese lenguaje: calcula, guarda tablas y dibuja gráficos. No hace falta conocerlo para explicar el proyecto, y no es una conclusión científica."),
        ("¿Qué es una semilla?", "El número que arranca el azar. Semilla 1 repetida da otra vez el recorrido de la semilla 1."),
        ("¿Por qué 40 por material?", "Para ver la media y la dispersión. Se había dejado abierta la duda entre 30 y 40. La campaña final usó 40."),
        ("¿Por qué el GO se mueve y el AC no?", "Es una decisión de este modelo: el GO es móvil y el AC es fijo. No es una filmación del material."),
        ("¿Por qué los tamaños son distintos?", "Para representar de forma distinta el colorante, el GO y el AC en el plano. Son píxeles de dibujo."),
        ("¿Qué es qmax?", "El techo de la curva de Langmuir de referencia. No es el resultado de una semilla."),
        ("¿Por qué qt puede pasar de 380,7?", "Porque el programa adsorbe objetos enteros y no está obligado a parar en el equilibrio de referencia. El máximo de la campaña en GO es 392 mg/g."),
        ("¿Qué significa P_ads = 1?", "Si hay contacto y aún cabe 0,02 mg, el objeto se adsorbe. Es la regla de la campaña, no una probabilidad medida."),
        ("¿Qué significa S* = 0,55?", "Un valor publicado para un carbón activo y este colorante. Se guardó como contexto. No se usó como regla de la campaña."),
        ("¿Por qué hay tantos CSV?", "Uno por corrida, más las tablas resumen. Así se puede abrir en una hoja de cálculo tanto el detalle como la media."),
        ("¿Qué es tests?", "Comprobaciones del programa. 135 en este repositorio. No son experimentos."),
        ("¿Qué es scripts?", "Herramientas auxiliares. El motor está en src, no en scripts."),
        ("¿Cómo sé que se puede repetir?", "Misma semilla, mismo material, mismo salto, mismos pasos y la misma regla. Eso está guardado junto al resultado."),
        ("¿Cómo se sabe que no se pierde masa?", "En las 80 corridas finales, lo adsorbido más lo que sigue libre recupera los 4 mg iniciales."),
        ("¿Más contactos implican más adsorción?", "No automáticamente. El GO contacta más. El AC convierte mejor cada contacto, porque muchos roces no encuentran sitio para otro objeto entero."),
    ]
    for q, a in pairs:
        pdf.qa(q, a)

    pdf.h1("18. Diez cosas que hay que saber sí o sí")
    ten = [
        "Es un modelo de objetos en un plano, no de moléculas.",
        "El ciclo es: salto, contacto, adsorción, qt.",
        "Se compara GO móvil con AC fijo, frente al azul de metileno.",
        "Condiciones: 25 °C, pH 6, 4 mg, 10 mg, 40 mL. C0 = 4 mg / 0,04 L = 100 mg/L, magnitud derivada.",
        "Campaña: salto 1,5, 2000 pasos, regla de adsorción 1, 40 semillas por material.",
        "GO: unos 187 objetos, 94 %, 375 mg/g. AC: unos 133, 67 %, 266 mg/g.",
        "El GO contacta más. El AC aprovecha mejor cada contacto.",
        "Langmuir y la curva en minutos están al lado. No mandan en el motor.",
        "Los tests comprueban el programa. Los CSV y JSON guardan los resultados.",
        "No decir que el laboratorio ya ha confirmado que el GO es mejor.",
    ]
    for i, text in enumerate(ten, 1):
        pdf.bullet(f"{i}. {text}")

    pdf.h1("19. Resumen para memorizar")
    pdf.p(
        "Agua con colorante azul, y un sólido que puede retenerlo: GO o AC. El programa mueve objetos al azar, mira si se tocan y, si cabe, adsorbe. Repetido 40 veces, el GO retiene más que el AC en este cálculo: unos 187 frente a unos 133, unos 375 mg/g frente a unos 266. La masa cuadra. La misma semilla se puede repetir. No es un experimento y no es dinámica molecular."
    )
    return pdf


def write_docx(blocks: list[tuple[str, str]]) -> str | None:
    try:
        from docx import Document
    except ImportError:
        return None
    doc = Document()
    doc.add_heading("Cómo explicar este proyecto", 0)
    doc.add_paragraph(
        "Guía de estudio para quien ha seguido el trabajo y no necesita programar para contarlo."
    )
    for kind, text in blocks:
        if kind == "h1":
            doc.add_heading(text, level=1)
        elif kind == "term":
            name, _, body = text.partition("\n")
            para = doc.add_paragraph()
            para.add_run(name).bold = True
            doc.add_paragraph(body)
        elif kind == "qa":
            q, _, a = text.partition("\n")
            para = doc.add_paragraph()
            para.add_run(q).bold = True
            doc.add_paragraph(a)
        elif kind == "b":
            doc.add_paragraph(text, style="List Bullet")
        else:
            doc.add_paragraph(text)
    path = DOCS / "guia_explicacion_proyecto.docx"
    doc.save(str(path))
    return str(path)


def main() -> None:
    DOCS.mkdir(parents=True, exist_ok=True)
    pdf = build()
    path = DOCS / "guia_explicacion_proyecto.pdf"
    pdf.output(str(path))
    print(f"pdf_paginas={pdf.page_no()} path={path}")
    docx_path = write_docx(pdf.blocks)
    print(f"docx={docx_path or 'no_disponible'}")


if __name__ == "__main__":
    main()
